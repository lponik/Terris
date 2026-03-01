#!/usr/bin/env python3
"""Build processed PFAS scoring datasets from manually downloaded raw files.

This script is offline-only and memory-safe for large FRS inputs.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import re
import sys
from collections.abc import Iterable
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

from process_superfund_npl import find_superfund_gdb, process_superfund_gdb

REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = REPO_ROOT / "data" / "raw"
PROCESSED_DIR = REPO_ROOT / "data" / "processed"

UNIFIED_COLUMNS = [
    "id",
    "name",
    "category",
    "lat",
    "lon",
    "state",
    "source",
    "metadata_json",
]

CATEGORY_VALUES = {
    "industrial_frs",
    "military_base",
    "landfill",
    "superfund_npl",
}

OUTPUT_FILES = {
    "industrial_frs": PROCESSED_DIR / "industrial_frs.csv",
    "military_base": PROCESSED_DIR / "military_base.csv",
    "landfill": PROCESSED_DIR / "landfill.csv",
    "superfund_npl": PROCESSED_DIR / "superfund_npl.csv",
    "all_sites": PROCESSED_DIR / "all_sites.csv",
}

PROCESS_STATS_PATH = PROCESSED_DIR / "process_stats.json"

FRS_COLUMN_CANDIDATES = {
    "registry_id": ["REGISTRY_ID", "REGISTRYID"],
    "name": ["PRIMARY_NAME", "PRIMARYNAME", "FACILITY_NAME"],
    "lat": ["LATITUDE83", "LATITUDE"],
    "lon": ["LONGITUDE83", "LONGITUDE"],
    "state": ["STATE_CODE", "STATE"],
}

MILITARY_COLUMN_CANDIDATES = {
    "id": [
        "GLOBALLY_UNIQUE_IDENTIFIER",
        "GLOBALLY UNIQUE IDENTIFIER",
        "PRIMARY_KEY_IDENTIFIER",
        "PRIMARY KEY IDENTIFIER",
        "OBJECTID",
        "SITE_ID",
        "ID",
    ],
    "name": [
        "SITE_NAME",
        "SITE NAME",
        "FEATURE_NAME",
        "FEATURE NAME",
        "NAME",
        "BASE_NAME",
        "BASE NAME",
        "INSTALLATION_NAME",
    ],
    "state": [
        "STATE",
        "STATE_CODE",
        "STATE_NAME_CODE",
        "ST",
        "STATE_ABBR",
    ],
    "lat": [
        "LATITUDE",
        "LAT",
        "DEC_LAT",
        "POINT_Y",
        "Y",
        "YCOORD",
        "SHAPE__Y",
    ],
    "lon": [
        "LONGITUDE",
        "LON",
        "LONG",
        "DEC_LON",
        "POINT_X",
        "X",
        "XCOORD",
        "SHAPE__X",
    ],
}

LANDFILL_COLUMN_CANDIDATES = {
    "id": ["LANDFILL_ID", "LANDFILL ID", "ID"],
    "name": ["LANDFILL_NAME", "LANDFILL NAME", "NAME"],
    "state": ["STATE", "STATE_CODE", "ST"],
    "lat": ["LATITUDE", "LAT"],
    "lon": ["LONGITUDE", "LON", "LONG"],
    "ghgrp_id": ["GHGRP_ID", "GHGRP ID"],
    "city": ["CITY"],
    "county": ["COUNTY"],
    "zip": ["ZIP_CODE", "ZIP CODE", "ZIP"],
    "address": ["PHYSICAL_ADDRESS", "PHYSICAL ADDRESS", "ADDRESS"],
}

MILITARY_NAME_REPLACEMENTS = {
    "tng": "training",
    "ctr": "center",
    "ctrs": "centers",
    "stn": "station",
    "sta": "station",
    "ft": "fort",
    "afb": "airforcebase",
    "ang": "airnationalguard",
    "ng": "nationalguard",
    "jrb": "jointreservebase",
    "jb": "jointbase",
}

MILITARY_NAME_STOPWORDS = {
    "the",
    "us",
    "usa",
    "u",
    "s",
    "united",
    "states",
    "military",
    "base",
    "site",
    "installation",
    "facility",
    "area",
    "and",
    "of",
    "at",
    "for",
    "forces",
}


def normalize_col(name: str | None) -> str:
    if not name:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(name).strip().lower())


def normalize_site_name(name: Any, canonical: bool = False) -> str:
    if name is None:
        return ""
    text = str(name).strip().lower().replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    tokens: list[str] = []
    for token in text.split():
        if canonical:
            token = MILITARY_NAME_REPLACEMENTS.get(token, token)
            if token in MILITARY_NAME_STOPWORDS:
                continue
        tokens.append(token)
    return " ".join(tokens)


def pick_column(fieldnames: Iterable[str], candidates: list[str]) -> str | None:
    lookup = {normalize_col(field): field for field in fieldnames if field}
    for candidate in candidates:
        key = normalize_col(candidate)
        if key in lookup:
            return lookup[key]
    return None


def parse_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def is_valid_coord(lat: float | None, lon: float | None) -> bool:
    if lat is None or lon is None:
        return False
    return -90.0 <= lat <= 90.0 and -180.0 <= lon <= 180.0


def clean_state(value: Any) -> str:
    if value is None:
        return ""
    state = str(value).strip().upper()
    if len(state) == 2 and state.isalpha():
        return state
    return ""


def clean_id_fragment(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text:
        return ""
    return re.sub(r"[^A-Za-z0-9_.-]+", "", text)


def stable_hash(*parts: str) -> str:
    joined = "|".join(part.strip() for part in parts if part is not None)
    return hashlib.sha1(joined.encode("utf-8")).hexdigest()[:16]


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_output_header(path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=UNIFIED_COLUMNS)
        writer.writeheader()


def init_csv_limits() -> None:
    max_int = sys.maxsize
    while True:
        try:
            csv.field_size_limit(max_int)
            return
        except OverflowError:
            max_int = int(max_int / 10)


def detect_raw_inputs(raw_dir: Path) -> dict[str, Path | None]:
    if not raw_dir.exists():
        raise FileNotFoundError(f"Missing raw data directory: {raw_dir}")

    files = [path for path in raw_dir.iterdir() if path.is_file()]

    def choose(
        *,
        must_contain_any: list[str],
        extensions: tuple[str, ...],
        prefer_extensions: tuple[str, ...] | None = None,
    ) -> Path | None:
        candidates = []
        for path in files:
            name = path.name.lower()
            if path.suffix.lower() not in extensions:
                continue
            if not any(token in name for token in must_contain_any):
                continue
            candidates.append(path)

        if not candidates:
            return None

        if prefer_extensions:
            ext_rank = {ext: idx for idx, ext in enumerate(prefer_extensions)}
            candidates.sort(
                key=lambda p: (ext_rank.get(p.suffix.lower(), 99), -p.stat().st_size, p.name)
            )
        else:
            candidates.sort(key=lambda p: (-p.stat().st_size, p.name))

        return candidates[0]

    frs_facility = choose(
        must_contain_any=["national_facility", "facility_file"],
        extensions=(".csv",),
    )

    if frs_facility is None:
        frs_facility = choose(
            must_contain_any=["facility", "frs"],
            extensions=(".csv",),
        )

    frs_program = choose(
        must_contain_any=["national_program", "program_file"],
        extensions=(".csv",),
    )

    military = choose(
        must_contain_any=["military", "base", "ntad"],
        extensions=(".csv", ".geojson", ".json"),
        prefer_extensions=(".geojson", ".json", ".csv"),
    )

    landfill = choose(
        must_contain_any=["lmop", "landfill"],
        extensions=(".xlsx",),
    )

    if frs_facility is None:
        raise FileNotFoundError(
            "Could not detect FRS facility CSV in data/raw. "
            "Expected a file like NATIONAL_FACILITY_FILE.CSV."
        )
    if military is None:
        raise FileNotFoundError(
            "Could not detect military bases file in data/raw "
            "(.csv, .geojson, or .json)."
        )
    if landfill is None:
        raise FileNotFoundError(
            "Could not detect LMOP landfill workbook (.xlsx) in data/raw."
        )

    return {
        "frs_facility": frs_facility,
        "frs_program": frs_program,
        "military": military,
        "landfill": landfill,
    }


def process_frs(
    source_path: Path,
    output_path: Path,
    progress_every: int,
) -> dict[str, int]:
    logging.info("Processing FRS facilities: %s", source_path.name)
    write_output_header(output_path)

    stats = {
        "input_rows": 0,
        "written_rows": 0,
        "dropped_rows": 0,
        "dropped_missing_required": 0,
        "dropped_invalid_coords": 0,
    }

    with source_path.open("r", encoding="utf-8", errors="replace", newline="") as source_handle:
        reader = csv.DictReader(source_handle)
        if not reader.fieldnames:
            raise ValueError("FRS CSV has no header row.")

        col_registry = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["registry_id"])
        col_name = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["name"])
        col_lat = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["lat"])
        col_lon = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["lon"])
        col_state = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["state"])

        required_mapping = {
            "registry_id": col_registry,
            "name": col_name,
            "lat": col_lat,
            "lon": col_lon,
        }
        missing = [name for name, column in required_mapping.items() if column is None]
        if missing:
            raise ValueError(
                f"FRS CSV missing required columns: {missing}. "
                f"Detected fields include: {reader.fieldnames[:20]}..."
            )

        metadata_columns = {
            "city": pick_column(reader.fieldnames, ["CITY_NAME", "CITY"]),
            "county": pick_column(reader.fieldnames, ["COUNTY_NAME", "COUNTY"]),
            "postal_code": pick_column(reader.fieldnames, ["POSTAL_CODE", "ZIP_CODE", "ZIP"]),
        }

        with output_path.open("a", encoding="utf-8", newline="") as output_handle:
            writer = csv.DictWriter(output_handle, fieldnames=UNIFIED_COLUMNS)

            for row in reader:
                stats["input_rows"] += 1
                if stats["input_rows"] % progress_every == 0:
                    logging.info(
                        "FRS progress rows=%s written=%s dropped=%s",
                        f"{stats['input_rows']:,}",
                        f"{stats['written_rows']:,}",
                        f"{stats['dropped_rows']:,}",
                    )

                registry_id = clean_id_fragment(row.get(col_registry))
                facility_name = str(row.get(col_name, "") or "").strip()
                lat = parse_float(row.get(col_lat))
                lon = parse_float(row.get(col_lon))

                if not registry_id or not facility_name or lat is None or lon is None:
                    stats["dropped_rows"] += 1
                    stats["dropped_missing_required"] += 1
                    continue

                if not is_valid_coord(lat, lon):
                    stats["dropped_rows"] += 1
                    stats["dropped_invalid_coords"] += 1
                    continue

                metadata: dict[str, str] = {}
                for key, col_name_opt in metadata_columns.items():
                    if col_name_opt is None:
                        continue
                    value = str(row.get(col_name_opt, "") or "").strip()
                    if value:
                        metadata[key] = value

                writer.writerow(
                    {
                        "id": f"frs_{registry_id}",
                        "name": facility_name,
                        "category": "industrial_frs",
                        "lat": f"{lat:.8f}",
                        "lon": f"{lon:.8f}",
                        "state": clean_state(row.get(col_state)) if col_state else "",
                        "source": "EPA_FRS_NATIONAL_FACILITY",
                        "metadata_json": json.dumps(metadata, separators=(",", ":"), ensure_ascii=True),
                    }
                )
                stats["written_rows"] += 1

    logging.info(
        "FRS done rows=%s written=%s dropped=%s",
        f"{stats['input_rows']:,}",
        f"{stats['written_rows']:,}",
        f"{stats['dropped_rows']:,}",
    )
    return stats


def geometry_centroid(geometry: dict[str, Any] | None) -> tuple[float | None, float | None]:
    if not geometry:
        return None, None

    coords = geometry.get("coordinates")
    if coords is None:
        return None, None

    points: list[tuple[float, float]] = []

    def collect(value: Any) -> None:
        if isinstance(value, (list, tuple)):
            if len(value) >= 2 and all(isinstance(value[idx], (int, float)) for idx in (0, 1)):
                points.append((float(value[0]), float(value[1])))
            else:
                for item in value:
                    collect(item)

    collect(coords)

    if not points:
        return None, None

    lon = sum(point[0] for point in points) / len(points)
    lat = sum(point[1] for point in points) / len(points)
    return lat, lon


def backfill_military_coords_from_frs(
    military_entries: list[dict[str, Any]],
    frs_source_path: Path,
    progress_every: int = 500_000,
) -> int:
    pending_state_basic: dict[tuple[str, str], list[int]] = {}
    pending_state_canonical: dict[tuple[str, str], list[int]] = {}
    pending_name_basic: dict[str, list[int]] = {}
    pending_name_canonical: dict[str, list[int]] = {}

    for idx, entry in enumerate(military_entries):
        if entry.get("lat") is not None and entry.get("lon") is not None:
            continue
        name = entry.get("name", "")
        state = entry.get("state", "")
        basic = normalize_site_name(name, canonical=False)
        canonical = normalize_site_name(name, canonical=True)
        if not basic and not canonical:
            continue

        if state:
            if basic:
                pending_state_basic.setdefault((state, basic), []).append(idx)
            if canonical:
                pending_state_canonical.setdefault((state, canonical), []).append(idx)
        else:
            if basic:
                pending_name_basic.setdefault(basic, []).append(idx)
            if canonical:
                pending_name_canonical.setdefault(canonical, []).append(idx)

    if not (
        pending_state_basic
        or pending_state_canonical
        or pending_name_basic
        or pending_name_canonical
    ):
        return 0

    matched_entries = 0
    scanned_rows = 0

    with frs_source_path.open("r", encoding="utf-8", errors="replace", newline="") as source_handle:
        reader = csv.DictReader(source_handle)
        if not reader.fieldnames:
            raise ValueError("FRS CSV has no header row for military coordinate backfill.")

        col_registry = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["registry_id"])
        col_name = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["name"])
        col_lat = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["lat"])
        col_lon = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["lon"])
        col_state = pick_column(reader.fieldnames, FRS_COLUMN_CANDIDATES["state"])

        required_mapping = {
            "registry_id": col_registry,
            "name": col_name,
            "lat": col_lat,
            "lon": col_lon,
        }
        missing = [name for name, column in required_mapping.items() if column is None]
        if missing:
            raise ValueError(
                f"FRS CSV missing required backfill columns: {missing}. "
                f"Detected fields include: {reader.fieldnames[:20]}..."
            )

        for row in reader:
            scanned_rows += 1
            if scanned_rows % progress_every == 0:
                logging.info(
                    "Military backfill progress rows=%s matched=%s",
                    f"{scanned_rows:,}",
                    f"{matched_entries:,}",
                )

            lat = parse_float(row.get(col_lat))
            lon = parse_float(row.get(col_lon))
            if not is_valid_coord(lat, lon):
                continue

            frs_name = str(row.get(col_name, "") or "").strip()
            if not frs_name:
                continue
            frs_state = clean_state(row.get(col_state)) if col_state else ""
            name_basic = normalize_site_name(frs_name, canonical=False)
            name_canonical = normalize_site_name(frs_name, canonical=True)
            registry_id = clean_id_fragment(row.get(col_registry))

            def assign(target_indexes: list[int], method: str) -> int:
                assigned = 0
                for target_idx in target_indexes:
                    target = military_entries[target_idx]
                    if target.get("lat") is not None and target.get("lon") is not None:
                        continue
                    target["lat"] = lat
                    target["lon"] = lon
                    metadata: dict[str, str] = target["metadata"]
                    metadata["coord_match_method"] = method
                    metadata["matched_frs_name"] = frs_name
                    if registry_id:
                        metadata["matched_frs_registry_id"] = registry_id
                    assigned += 1
                return assigned

            if frs_state and name_basic:
                matched_entries += assign(
                    pending_state_basic.get((frs_state, name_basic), []),
                    "frs_state_name_exact",
                )

            if frs_state and name_canonical:
                matched_entries += assign(
                    pending_state_canonical.get((frs_state, name_canonical), []),
                    "frs_state_name_canonical",
                )

            if name_basic:
                matched_entries += assign(
                    pending_name_basic.get(name_basic, []),
                    "frs_name_exact",
                )

            if name_canonical:
                matched_entries += assign(
                    pending_name_canonical.get(name_canonical, []),
                    "frs_name_canonical",
                )

    logging.info(
        "Military coordinate backfill complete scanned=%s matched=%s",
        f"{scanned_rows:,}",
        f"{matched_entries:,}",
    )
    return matched_entries


def process_military_csv(
    source_path: Path,
    output_path: Path,
    frs_source_path: Path | None = None,
) -> dict[str, int]:
    logging.info("Processing military CSV: %s", source_path.name)
    write_output_header(output_path)

    stats = {
        "input_rows": 0,
        "written_rows": 0,
        "dropped_rows": 0,
        "dropped_missing_required": 0,
        "dropped_invalid_coords": 0,
    }

    military_entries: list[dict[str, Any]] = []

    with source_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source_handle:
        reader = csv.DictReader(source_handle)
        if not reader.fieldnames:
            raise ValueError("Military CSV has no header row.")

        col_id = pick_column(reader.fieldnames, MILITARY_COLUMN_CANDIDATES["id"])
        col_name = pick_column(reader.fieldnames, MILITARY_COLUMN_CANDIDATES["name"])
        col_state = pick_column(reader.fieldnames, MILITARY_COLUMN_CANDIDATES["state"])
        col_lat = pick_column(reader.fieldnames, MILITARY_COLUMN_CANDIDATES["lat"])
        col_lon = pick_column(reader.fieldnames, MILITARY_COLUMN_CANDIDATES["lon"])

        extra_metadata_columns = {
            "site_operational_status": pick_column(
                reader.fieldnames,
                ["SITE_OPERATIONAL_STATUS", "SITE OPERATIONAL STATUS"],
            ),
            "reporting_component": pick_column(
                reader.fieldnames,
                ["SITE_REPORTING_COMPONENT_CODE", "SITE REPORTING COMPONENT CODE"],
            ),
        }

        for row in reader:
            stats["input_rows"] += 1
            name = str(row.get(col_name, "") or "").strip() if col_name else ""
            state = clean_state(row.get(col_state)) if col_state else ""
            stable_id = clean_id_fragment(row.get(col_id)) if col_id else ""

            lat = parse_float(row.get(col_lat)) if col_lat else None
            lon = parse_float(row.get(col_lon)) if col_lon else None

            metadata: dict[str, str] = {}
            for key, col_name_opt in extra_metadata_columns.items():
                if col_name_opt is None:
                    continue
                value = str(row.get(col_name_opt, "") or "").strip()
                if value:
                    metadata[key] = value

            military_entries.append(
                {
                    "row_num": stats["input_rows"] + 1,
                    "stable_id": stable_id,
                    "name": name,
                    "state": state,
                    "lat": lat,
                    "lon": lon,
                    "metadata": metadata,
                }
            )

    has_coord_columns = col_lat is not None and col_lon is not None
    if not has_coord_columns:
        if frs_source_path is None:
            logging.warning(
                "Military CSV has no coordinate columns and FRS backfill source was not provided."
            )
        else:
            logging.warning(
                "Military CSV has no coordinate columns. "
                "Attempting offline coordinate backfill from FRS facility names."
            )
            backfill_military_coords_from_frs(
                military_entries=military_entries,
                frs_source_path=frs_source_path,
            )

    with output_path.open("a", encoding="utf-8", newline="") as output_handle:
        writer = csv.DictWriter(output_handle, fieldnames=UNIFIED_COLUMNS)

        for entry in military_entries:
            lat = entry.get("lat")
            lon = entry.get("lon")
            if lat is None or lon is None:
                stats["dropped_rows"] += 1
                stats["dropped_missing_required"] += 1
                continue

            if not is_valid_coord(lat, lon):
                stats["dropped_rows"] += 1
                stats["dropped_invalid_coords"] += 1
                continue

            stable_id = entry.get("stable_id") or stable_hash(
                entry.get("name", ""),
                f"{lat:.8f}",
                f"{lon:.8f}",
            )

            writer.writerow(
                {
                    "id": f"mil_{stable_id}",
                    "name": entry.get("name") or f"Military site {stable_id}",
                    "category": "military_base",
                    "lat": f"{lat:.8f}",
                    "lon": f"{lon:.8f}",
                    "state": entry.get("state", ""),
                    "source": "BTS_MIL_BASES",
                    "metadata_json": json.dumps(
                        entry["metadata"],
                        separators=(",", ":"),
                        ensure_ascii=True,
                    ),
                }
            )
            stats["written_rows"] += 1

    if stats["written_rows"] == 0:
        logging.warning(
            "Military CSV produced 0 rows with valid coordinates. "
            "If this source is polygon-only, provide GeoJSON with geometry or a CSV with lat/lon fields."
        )

    logging.info(
        "Military CSV done rows=%s written=%s dropped=%s",
        f"{stats['input_rows']:,}",
        f"{stats['written_rows']:,}",
        f"{stats['dropped_rows']:,}",
    )
    return stats


def process_military_geojson(source_path: Path, output_path: Path) -> dict[str, int]:
    logging.info("Processing military GeoJSON/JSON: %s", source_path.name)
    write_output_header(output_path)

    stats = {
        "input_rows": 0,
        "written_rows": 0,
        "dropped_rows": 0,
        "dropped_missing_required": 0,
        "dropped_invalid_coords": 0,
    }

    with source_path.open("r", encoding="utf-8", errors="replace") as handle:
        payload = json.load(handle)

    if isinstance(payload, dict) and isinstance(payload.get("features"), list):
        features = payload["features"]
    elif isinstance(payload, list):
        features = payload
    else:
        raise ValueError("Military JSON does not look like a FeatureCollection or feature list.")

    with output_path.open("a", encoding="utf-8", newline="") as output_handle:
        writer = csv.DictWriter(output_handle, fieldnames=UNIFIED_COLUMNS)

        for feature in features:
            stats["input_rows"] += 1
            properties = feature.get("properties") if isinstance(feature, dict) else {}
            if not isinstance(properties, dict):
                properties = {}

            lat, lon = geometry_centroid(feature.get("geometry") if isinstance(feature, dict) else None)
            if lat is None or lon is None:
                stats["dropped_rows"] += 1
                stats["dropped_missing_required"] += 1
                continue

            if not is_valid_coord(lat, lon):
                stats["dropped_rows"] += 1
                stats["dropped_invalid_coords"] += 1
                continue

            name = ""
            for candidate in MILITARY_COLUMN_CANDIDATES["name"]:
                for key in properties:
                    if normalize_col(key) == normalize_col(candidate):
                        name = str(properties.get(key) or "").strip()
                        break
                if name:
                    break

            state = ""
            for candidate in MILITARY_COLUMN_CANDIDATES["state"]:
                for key in properties:
                    if normalize_col(key) == normalize_col(candidate):
                        state = clean_state(properties.get(key))
                        break
                if state:
                    break

            stable_id = ""
            for candidate in MILITARY_COLUMN_CANDIDATES["id"]:
                for key in properties:
                    if normalize_col(key) == normalize_col(candidate):
                        stable_id = clean_id_fragment(properties.get(key))
                        break
                if stable_id:
                    break

            if not stable_id:
                stable_id = stable_hash(name, f"{lat:.8f}", f"{lon:.8f}")

            metadata: dict[str, str] = {}
            for key in ("site_operational_status", "site_reporting_component_code"):
                for prop_key, prop_value in properties.items():
                    if normalize_col(prop_key) == normalize_col(key) and prop_value not in (None, ""):
                        metadata[key] = str(prop_value).strip()

            writer.writerow(
                {
                    "id": f"mil_{stable_id}",
                    "name": name or f"Military site {stable_id}",
                    "category": "military_base",
                    "lat": f"{lat:.8f}",
                    "lon": f"{lon:.8f}",
                    "state": state,
                    "source": "BTS_MIL_BASES",
                    "metadata_json": json.dumps(metadata, separators=(",", ":"), ensure_ascii=True),
                }
            )
            stats["written_rows"] += 1

    logging.info(
        "Military GeoJSON done rows=%s written=%s dropped=%s",
        f"{stats['input_rows']:,}",
        f"{stats['written_rows']:,}",
        f"{stats['dropped_rows']:,}",
    )
    return stats


def detect_landfill_sheet(workbook: Any) -> Any:
    prioritized = sorted(
        workbook.worksheets,
        key=lambda ws: (0 if "database" in ws.title.lower() else 1, ws.title.lower()),
    )

    for worksheet in prioritized:
        header_row = find_header_row(worksheet)
        if header_row is not None:
            return worksheet, header_row

    raise ValueError("Could not find a landfill worksheet with latitude/longitude columns.")


def find_header_row(worksheet: Any, max_scan_rows: int = 40) -> int | None:
    for row_index, row_values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        normalized_values = {normalize_col(value): idx for idx, value in enumerate(row_values) if value}
        has_lat = any("latitude" in key for key in normalized_values)
        has_lon = any("longitude" in key for key in normalized_values)
        has_name = any("landfill" in key and "name" in key for key in normalized_values)
        if has_lat and has_lon and has_name:
            return row_index
    return None


def process_landfill_xlsx(source_path: Path, output_path: Path) -> dict[str, int]:
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is required for landfill processing. Install dependencies first "
            "(for example: pip install -r requirements.txt)."
        )

    logging.info("Processing landfill workbook: %s", source_path.name)
    write_output_header(output_path)

    stats = {
        "input_rows": 0,
        "written_rows": 0,
        "dropped_rows": 0,
        "dropped_missing_required": 0,
        "dropped_invalid_coords": 0,
    }

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet, header_row = detect_landfill_sheet(workbook)

    header_values = next(
        worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        )
    )
    headers = [str(value).strip() if value is not None else "" for value in header_values]

    def header_index(candidates: list[str]) -> int | None:
        chosen = pick_column(headers, candidates)
        if chosen is None:
            return None
        return headers.index(chosen)

    idx_id = header_index(LANDFILL_COLUMN_CANDIDATES["id"])
    idx_name = header_index(LANDFILL_COLUMN_CANDIDATES["name"])
    idx_state = header_index(LANDFILL_COLUMN_CANDIDATES["state"])
    idx_lat = header_index(LANDFILL_COLUMN_CANDIDATES["lat"])
    idx_lon = header_index(LANDFILL_COLUMN_CANDIDATES["lon"])

    required = {
        "name": idx_name,
        "lat": idx_lat,
        "lon": idx_lon,
    }
    missing = [key for key, idx in required.items() if idx is None]
    if missing:
        raise ValueError(
            f"Landfill worksheet '{worksheet.title}' missing required columns: {missing}."
        )

    metadata_indexes = {
        "ghgrp_id": header_index(LANDFILL_COLUMN_CANDIDATES["ghgrp_id"]),
        "city": header_index(LANDFILL_COLUMN_CANDIDATES["city"]),
        "county": header_index(LANDFILL_COLUMN_CANDIDATES["county"]),
        "zip": header_index(LANDFILL_COLUMN_CANDIDATES["zip"]),
        "address": header_index(LANDFILL_COLUMN_CANDIDATES["address"]),
    }

    with output_path.open("a", encoding="utf-8", newline="") as output_handle:
        writer = csv.DictWriter(output_handle, fieldnames=UNIFIED_COLUMNS)

        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            stats["input_rows"] += 1

            if not row:
                stats["dropped_rows"] += 1
                stats["dropped_missing_required"] += 1
                continue

            name_value = row[idx_name] if idx_name is not None and idx_name < len(row) else None
            lat_value = row[idx_lat] if idx_lat is not None and idx_lat < len(row) else None
            lon_value = row[idx_lon] if idx_lon is not None and idx_lon < len(row) else None

            name = str(name_value).strip() if name_value is not None else ""
            lat = parse_float(lat_value)
            lon = parse_float(lon_value)

            if not name or lat is None or lon is None:
                stats["dropped_rows"] += 1
                stats["dropped_missing_required"] += 1
                continue

            if not is_valid_coord(lat, lon):
                stats["dropped_rows"] += 1
                stats["dropped_invalid_coords"] += 1
                continue

            landfill_id = ""
            if idx_id is not None and idx_id < len(row):
                landfill_id = clean_id_fragment(row[idx_id])
            if not landfill_id:
                landfill_id = stable_hash(name, f"{lat:.8f}", f"{lon:.8f}")

            state_value = row[idx_state] if idx_state is not None and idx_state < len(row) else ""
            state = clean_state(state_value)

            metadata: dict[str, str] = {}
            for key, idx in metadata_indexes.items():
                if idx is None or idx >= len(row):
                    continue
                value = row[idx]
                if value is None:
                    continue
                text = str(value).strip()
                if text:
                    metadata[key] = text

            writer.writerow(
                {
                    "id": f"landfill_{landfill_id}",
                    "name": name,
                    "category": "landfill",
                    "lat": f"{lat:.8f}",
                    "lon": f"{lon:.8f}",
                    "state": state,
                    "source": "EPA_LMOP",
                    "metadata_json": json.dumps(metadata, separators=(",", ":"), ensure_ascii=True),
                }
            )
            stats["written_rows"] += 1

    logging.info(
        "Landfill done rows=%s written=%s dropped=%s (sheet=%s)",
        f"{stats['input_rows']:,}",
        f"{stats['written_rows']:,}",
        f"{stats['dropped_rows']:,}",
        worksheet.title,
    )
    return stats


def concatenate_outputs(output_paths: list[Path], destination: Path) -> int:
    write_output_header(destination)
    total_rows = 0

    with destination.open("a", encoding="utf-8", newline="") as destination_handle:
        writer = csv.DictWriter(destination_handle, fieldnames=UNIFIED_COLUMNS)

        for output_path in output_paths:
            with output_path.open("r", encoding="utf-8", newline="") as source_handle:
                reader = csv.DictReader(source_handle)
                for row in reader:
                    writer.writerow({column: row.get(column, "") for column in UNIFIED_COLUMNS})
                    total_rows += 1

    logging.info("Concatenated %s total rows into %s", f"{total_rows:,}", destination.name)
    return total_rows


def prepare_outputs(overwrite: bool) -> None:
    ensure_dir(PROCESSED_DIR)
    targets = list(OUTPUT_FILES.values()) + [PROCESS_STATS_PATH]

    if overwrite:
        for target in targets:
            if target.exists():
                target.unlink()
        return

    existing = [path for path in targets if path.exists()]
    if existing:
        formatted = ", ".join(path.name for path in existing)
        raise FileExistsError(
            "Processed outputs already exist. Re-run with --overwrite to regenerate cleanly. "
            f"Existing: {formatted}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Process manual raw datasets into offline PFAS scoring bundle files."
        )
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Delete and regenerate files in data/processed.",
    )
    parser.add_argument(
        "--frs-progress-every",
        type=int,
        default=100_000,
        help="Log FRS progress every N rows (default: 100000).",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    init_csv_limits()

    try:
        prepare_outputs(overwrite=args.overwrite)
        raw_inputs = detect_raw_inputs(RAW_DIR)
        raw_inputs["superfund_npl_gdb"] = find_superfund_gdb(RAW_DIR)

        logging.info("Detected raw files:")
        for key, value in raw_inputs.items():
            logging.info("  %s: %s", key, value.name if value else "<not found>")

        dataset_stats: dict[str, dict[str, int]] = {}
        dataset_stats["industrial_frs"] = process_frs(
            source_path=raw_inputs["frs_facility"],
            output_path=OUTPUT_FILES["industrial_frs"],
            progress_every=max(args.frs_progress_every, 1),
        )

        military_input = raw_inputs["military"]
        if military_input.suffix.lower() == ".csv":
            dataset_stats["military_base"] = process_military_csv(
                source_path=military_input,
                output_path=OUTPUT_FILES["military_base"],
                frs_source_path=raw_inputs["frs_facility"],
            )
        else:
            dataset_stats["military_base"] = process_military_geojson(
                source_path=military_input,
                output_path=OUTPUT_FILES["military_base"],
            )

        dataset_stats["landfill"] = process_landfill_xlsx(
            source_path=raw_inputs["landfill"],
            output_path=OUTPUT_FILES["landfill"],
        )

        superfund_source = raw_inputs.get("superfund_npl_gdb")
        if superfund_source is not None:
            dataset_stats["superfund_npl"] = process_superfund_gdb(
                source_path=superfund_source,
                output_path=OUTPUT_FILES["superfund_npl"],
            )
        else:
            logging.info("Superfund .gdb not found; skipping")

        concat_sources = [
            OUTPUT_FILES["industrial_frs"],
            OUTPUT_FILES["military_base"],
            OUTPUT_FILES["landfill"],
        ]
        if "superfund_npl" in dataset_stats:
            concat_sources.append(OUTPUT_FILES["superfund_npl"])

        all_sites_rows = concatenate_outputs(
            concat_sources,
            OUTPUT_FILES["all_sites"],
        )

        payload = {
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "raw_inputs": {
                key: (str(path.relative_to(REPO_ROOT)) if path is not None else None)
                for key, path in raw_inputs.items()
            },
            "datasets": dataset_stats,
            "all_sites_rows": all_sites_rows,
        }
        PROCESS_STATS_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

        logging.info("Wrote process stats: %s", PROCESS_STATS_PATH)
        logging.info("Processing complete.")
        return 0

    except Exception as exc:  # pragma: no cover
        logging.error("Processing failed: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
